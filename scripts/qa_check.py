# -*- coding: utf-8 -*-
"""学术 deck 交付前质检：出界 / 文字溢出 / 样板残留词 / 数字与语料交叉核验 / 版面。

前四道门对应真实返工：克隆版式后忘记收缩文本框、忘记替换样板词、
口述时说的数在页面上对不上语料。版面门（重叠 / 图未填满面板 / 空白区 /
Unicode 下标 / 窄框两端对齐）来自论坛 deck 实测，都是目检漏掉、用户一眼看出的问题。
版面门全是粗估，只报 ⚠️；PNG 目检由 export_png.ps1 负责。
"""

import argparse
import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN

EMU_PER_INCH = 914400.0
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
TOLERANCE_IN = 0.05
DEFAULT_SIZE_PT = 18.0          # 读不到 presentation.xml 的默认字号时才用

# 参考样板 deck 自带的词，出现即说明该页没替换干净。
# 两类：样板示例的学科词汇，和匿名化时填入的占位符。
LSTALE_DEFAULT = [
    "水凝胶", "粘接", "剥离", "搭接剪切", "脱粘", "PAAm", "PAA", "PDMA",
    "拓扑缠结", "速率效应", "界面失效模式一", "材料–基底界面",
    "创新成果一", "创新成果二", "创新成果三", "创新成果四",
    "张三", "李四", "20XX", "占位条目", "Placeholder", "Author A",
]
# 只核验"数据型"数字：带小数点，或三位以上整数。章节号 / (a) / 01 不进核验
RE_DATA_NUMBER = re.compile(r"\d+\.\d+|\d{3,}")
# 引文年份不是数据主张，排除；真当数据用的年份由 PNG 目检兜底
RE_YEAR = re.compile(r"^(19|20)\d{2}$")
# 图注 / 引文行本就允许文字略溢出标称框，不参与溢出估算
OVERFLOW_EXEMPT_PT = 12.0
# 版面门阈值（英寸 / 比例）
OVERLAP_MIN_FRAC = 0.2          # 重叠面积占较小框的比例超过它才报；并排框的标称边界常有细缝重叠，不算
FILL_MIN_W, FILL_MIN_H = 0.85, 0.80   # 图宽与图高都低于面板的这个比例 = 没填满；旗标 + 图注约占面板高 20%，被高度卡住的图只能靠收窄面板达标
PANEL_MIN_IN2 = 4.0             # 小于它的无字形状是装饰条 / 灰条，不当面板
EMPTY_MAX_COVER = 0.05          # 内容区 4×2 网格里，形状覆盖率低于它的格子算空白
CONTENT_TOP_IN, CONTENT_BOTTOM_IN = 0.9, 7.1   # 标题栏以下、页脚以上
DECOR_MIN_FRAC = 0.5            # 有形状盖住半页以上 = 封面 / 目录 / 致谢等装饰页，不查空白
JUSTIFY_MAX_WIDTH_IN = 3.5      # 窄于它的框里两端对齐会把中英混排拉稀
RE_UNICODE_SCRIPT = re.compile("[\u2070-\u209c\u1d62-\u1d6a]")   # ₀ₐ ⁰ᵢ 等；² ³ 常用于单位，不在此列
LALIGN_JUSTIFY = (PP_ALIGN.JUSTIFY, PP_ALIGN.DISTRIBUTE, PP_ALIGN.JUSTIFY_LOW,
                  PP_ALIGN.THAI_DISTRIBUTE)


def fail(msg):
    print("❌ ERROR: " + msg)
    raise SystemExit(1)


def warn(msg):
    print("⚠️  " + msg)


def check_deck_exists(path_deck):
    if not path_deck.is_file():
        fail("deck 不存在 " + str(path_deck))
    return path_deck


def get_shape_type(shape):
    """python-pptx 对 officecli 生成的 OMML 公式形状（无 prstGeom 的 sp）取 shape_type 会抛异常，当作未知类型。"""
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def get_shape_text(shape):
    """形状（含组合内嵌）的纯文本，用换行拼接。"""
    if get_shape_type(shape) == 6:  # msoGroup
        return "\n".join(get_shape_text(sub) for sub in shape.shapes)
    if not shape.has_text_frame:
        return ""
    return shape.text_frame.text


def get_default_size_pt(prs):
    """presentation.xml 的 defaultTextStyle 一级字号：文本框没写字号时继承的就是它。"""
    lsz = prs.part._element.xpath("./p:defaultTextStyle/a:lvl1pPr/a:defRPr/@sz")
    return int(lsz[0]) / 100.0 if lsz else DEFAULT_SIZE_PT


def get_first_size_pt(shape):
    """取形状内第一个显式字号；继承时返回 deck 默认字号，仅用于溢出与重叠估算。"""
    if not shape.has_text_frame:
        return DEFAULT_SIZE_PT
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                return run.font.size.pt
    return DEFAULT_SIZE_PT


def check_out_of_canvas(shape):
    """带文字的形状超出画布返回描述字符串，否则 None。

    模板的背景块、页码徽章故意出血到画布外，是设计而非缺陷，所以只查有文字的形状。
    """
    if not get_shape_text(shape).strip():
        return None
    try:
        left, top = shape.left / EMU_PER_INCH, shape.top / EMU_PER_INCH
        right = left + shape.width / EMU_PER_INCH
        bottom = top + shape.height / EMU_PER_INCH
    except TypeError:
        return None
    if (left < -TOLERANCE_IN or top < -TOLERANCE_IN
            or right > SLIDE_W_IN + TOLERANCE_IN or bottom > SLIDE_H_IN + TOLERANCE_IN):
        return "越界 (%.2f,%.2f)–(%.2f,%.2f)" % (left, top, right, bottom)
    return None


def get_bbox_in(shape):
    """(left, top, right, bottom) 英寸；无几何（占位符继承）时返回 None。"""
    try:
        left, top = shape.left / EMU_PER_INCH, shape.top / EMU_PER_INCH
        return left, top, left + shape.width / EMU_PER_INCH, top + shape.height / EMU_PER_INCH
    except TypeError:
        return None


def get_ems(text):
    return sum(1.0 if ord(ch) > 0x2E80 else 0.5 for ch in text)


def get_text_extent_in(shape):
    """按 CJK 一字一 em、西文半 em、段落行距估算文字实际渲染的 (宽, 高)；无文字返回 None。

    框是"适应文字"时 PowerPoint 按实际文字撑开高度，标称框高会低估；
    引文行这类不换行的长框，文字只占左边一小段，标称框宽会高估。重叠检查两头都要用估算值。
    """
    text = get_shape_text(shape).strip()
    if not text or get_shape_type(shape) == MSO_SHAPE_TYPE.GROUP or not shape.has_text_frame:
        return None
    bbox = get_bbox_in(shape)
    size_in = get_first_size_pt(shape) / 72.0
    if bbox is None or size_in <= 0 or bbox[2] - bbox[0] <= 0:
        return None
    width_box = bbox[2] - bbox[0]
    ems_per_line = width_box / size_in
    spacing = shape.text_frame.paragraphs[0].line_spacing
    spacing = spacing if isinstance(spacing, float) else 1.0   # 固定磅值行距少见，按单倍估
    lines, ems_max = 0, 0.0
    for raw in text.split("\n"):
        ems = get_ems(raw)
        ems_max = max(ems_max, ems)
        if shape.text_frame.word_wrap is False:
            lines += 1
        else:
            lines += max(1, math.ceil(ems / ems_per_line))
    width = ems_max * size_in + 0.2                   # 0.2 = 左右默认内边距
    if shape.text_frame.word_wrap is not False:
        width = min(width, width_box)
    return width, lines * size_in * 1.25 * spacing    # 1.25 = 行高 / 字号


def get_text_bbox_in(shape):
    """估算的文字渲染包围盒：宽按对齐方式收窄，高按估算行数撑开（只往下长）。"""
    bbox = get_bbox_in(shape)
    extent = get_text_extent_in(shape)
    if bbox is None or extent is None:
        return bbox
    width, height = extent
    align = shape.text_frame.paragraphs[0].alignment
    if align == PP_ALIGN.CENTER:
        left = (bbox[0] + bbox[2] - width) / 2
    elif align == PP_ALIGN.RIGHT:
        left = bbox[2] - width
    else:
        left = bbox[0]
    return left, bbox[1], left + width, max(bbox[3], bbox[1] + height)


def check_text_overflow(shape):
    """估算文字高度超出框高返回描述字符串。

    只是粗估，用于抓"克隆后塞了两倍文字"这种明显溢出，误报按 ⚠️ 处理不阻断。
    "适应文字"的框会自己长高，不算溢出；它长高后压到别人由重叠检查负责。
    """
    bbox = get_bbox_in(shape)
    extent = get_text_extent_in(shape)
    if (bbox is None or extent is None or get_first_size_pt(shape) <= OVERFLOW_EXEMPT_PT
            or shape.text_frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT):
        return None
    needed_in, height_in = extent[1], bbox[3] - bbox[1]
    if needed_in > height_in * 1.15:            # 留 15% 余量再报
        return "文字溢出 需%.2f in / 框高%.2f in" % (needed_in, height_in)
    return None


def get_area_in2(bbox):
    return max(0.0, bbox[2] - bbox[0]) * max(0.0, bbox[3] - bbox[1])


def get_intersection(bbox_a, bbox_b):
    return (max(bbox_a[0], bbox_b[0]), max(bbox_a[1], bbox_b[1]),
            min(bbox_a[2], bbox_b[2]), min(bbox_a[3], bbox_b[3]))


def get_label(shape):
    return get_shape_text(shape).strip().replace("\n", "/")[:12]


def check_text_overlap(lshape):
    """文字形状两两重叠（按估算的文字渲染包围盒），返回描述列表。"""
    ltext = []
    for shape in lshape:
        bbox = get_text_bbox_in(shape)
        if bbox is None or not get_shape_text(shape).strip():
            continue
        ltext.append((shape, bbox))
    lmsg = []
    for i, (shape_a, bbox_a) in enumerate(ltext):
        for shape_b, bbox_b in ltext[i + 1:]:
            area_min = min(get_area_in2(bbox_a), get_area_in2(bbox_b))
            area = get_area_in2(get_intersection(bbox_a, bbox_b))
            if area_min > 0 and area / area_min > OVERLAP_MIN_FRAC:
                lmsg.append("文字重叠「%s」×「%s」" % (get_label(shape_a), get_label(shape_b)))
    return lmsg


def get_lpanel(lshape):
    """无字的自选图形且面积够大 = 面板框（双面板页的圆角矩形、总结页的栏框）。"""
    lpanel = []
    for shape in lshape:
        bbox = get_bbox_in(shape)
        if (get_shape_type(shape) == MSO_SHAPE_TYPE.AUTO_SHAPE and bbox is not None
                and not get_shape_text(shape).strip()
                and PANEL_MIN_IN2 < get_area_in2(bbox) < SLIDE_W_IN * SLIDE_H_IN * DECOR_MIN_FRAC):
            lpanel.append(bbox)
    return lpanel


def check_picture_fill(lshape):
    """图片在它所在面板里宽、高都没撑满，返回描述列表。"""
    lpanel = get_lpanel(lshape)
    lmsg = []
    for shape in lshape:
        bbox = get_bbox_in(shape)
        if get_shape_type(shape) != MSO_SHAPE_TYPE.PICTURE or bbox is None:
            continue
        x, y = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
        lhost = [p for p in lpanel if p[0] <= x <= p[2] and p[1] <= y <= p[3]]
        if not lhost:
            continue
        panel = min(lhost, key=get_area_in2)
        ratio_w = (bbox[2] - bbox[0]) / (panel[2] - panel[0])
        ratio_h = (bbox[3] - bbox[1]) / (panel[3] - panel[1])
        if ratio_w < FILL_MIN_W and ratio_h < FILL_MIN_H:
            lmsg.append("图未填满面板 宽%.0f%% 高%.0f%%（面板左上 %.1f,%.1f）"
                        % (ratio_w * 100, ratio_h * 100, panel[0], panel[1]))
    return lmsg


def check_empty_region(lshape):
    """内容区切成 4×2 格，几乎没有形状的格子 = 空白区（如图阵页右下角）。装饰页跳过。"""
    lbbox = [b for b in (get_bbox_in(s) for s in lshape) if b is not None]
    if any(get_area_in2(b) > SLIDE_W_IN * SLIDE_H_IN * DECOR_MIN_FRAC for b in lbbox):
        return []
    if not any(get_shape_type(s) == MSO_SHAPE_TYPE.PICTURE for s in lshape):
        return []
    n_col, n_row = 4, 2
    w = SLIDE_W_IN / n_col
    h = (CONTENT_BOTTOM_IN - CONTENT_TOP_IN) / n_row
    lmsg = []
    for row in range(n_row):
        for col in range(n_col):
            cell = (col * w, CONTENT_TOP_IN + row * h, (col + 1) * w, CONTENT_TOP_IN + (row + 1) * h)
            # 重叠部分会重复计，只让结果偏大，不会误报空白
            cover = sum(get_area_in2(get_intersection(cell, b)) for b in lbbox) / get_area_in2(cell)
            if cover < EMPTY_MAX_COVER:
                lmsg.append("空白区 第%d行第%d列（x %.1f–%.1f in，y %.1f–%.1f in）"
                            % (row + 1, col + 1, cell[0], cell[2], cell[1], cell[3]))
    return lmsg


def check_typography(shape):
    """Unicode 上下标字符、窄框两端对齐，返回描述列表。"""
    if not shape.has_text_frame:
        return []
    lmsg = []
    lhit = sorted(set(RE_UNICODE_SCRIPT.findall(shape.text_frame.text)))
    if lhit:
        lmsg.append("Unicode 上下标 %s（「%s」）→ 改用 subscript/superscript run"
                    % ("".join(lhit), get_label(shape)))
    bbox = get_bbox_in(shape)
    if bbox is not None and bbox[2] - bbox[0] < JUSTIFY_MAX_WIDTH_IN and any(
            p.alignment in LALIGN_JUSTIFY for p in shape.text_frame.paragraphs):
        lmsg.append("窄框两端对齐「%s」→ 改左对齐" % get_label(shape))
    return lmsg


def check_stale_words(text, lstale):
    return [word for word in lstale if word in text]


def get_source_numbers(path_source):
    """语料里出现过的数据型数字集合；无语料时返回 None 表示跳过核验。"""
    if path_source is None:
        return None
    if not path_source.is_file():
        fail("语料文件不存在 " + str(path_source))
    return set(RE_DATA_NUMBER.findall(path_source.read_text(encoding="utf-8")))


def scan_slide(index, slide, lstale, set_source_number):
    """返回 (l错误, l警告)。"""
    lerror, lwarn = [], []
    ltext = []
    lshape = list(slide.shapes)
    for shape in lshape:
        msg = check_out_of_canvas(shape)
        if msg:
            lerror.append(msg)
        msg = check_text_overflow(shape)
        if msg:
            lwarn.append(msg)
        lwarn += check_typography(shape)
        ltext.append(get_shape_text(shape))
    text = "\n".join(ltext)
    lwarn += check_text_overlap(lshape) + check_picture_fill(lshape) + check_empty_region(lshape)

    lhit = check_stale_words(text, lstale)
    if lhit:
        lerror.append("样板残留词 " + "、".join(lhit))

    if set_source_number is not None:
        lfound = {n for n in RE_DATA_NUMBER.findall(text) if not RE_YEAR.match(n)}
        lorphan = sorted(lfound - set_source_number)
        if lorphan:
            lerror.append("语料中查无此数 " + "、".join(lorphan))

    mark = "❌" if lerror else ("⚠️ " if lwarn else "✅")
    print("  %s 第 %d 页" % (mark, index))
    for msg in lerror + lwarn:
        print("       " + msg)
    return lerror, lwarn


def build_parser():
    parser = argparse.ArgumentParser(description="学术 deck 交付前质检")
    parser.add_argument("-deck", required=True, help="待检 .pptx")
    parser.add_argument("-source", default=None,
                        help="内容语料文本；给了才做数字交叉核验")
    parser.add_argument("-stale", default=None,
                        help="自定义残留词，逗号分隔；默认查参考样板 deck 的词")
    return parser


args = build_parser().parse_args()
sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)

path_deck = check_deck_exists(Path(args.deck).resolve())
path_source = Path(args.source).resolve() if args.source else None
lstale = [w.strip() for w in args.stale.split(",")] if args.stale else LSTALE_DEFAULT
set_source_number = get_source_numbers(path_source)

print("📁 %s" % path_deck.name)
print("📍 残留词 %d 个 ｜ 数字核验 %s"
      % (len(lstale), "开" if set_source_number is not None else "关（未给 -source）"))

prs = Presentation(str(path_deck))
DEFAULT_SIZE_PT = get_default_size_pt(prs)
n_error, n_warn = 0, 0
for index, slide in enumerate(prs.slides, 1):
    lerror, lwarn = scan_slide(index, slide, lstale, set_source_number)
    n_error += len(lerror)
    n_warn += len(lwarn)

print("📊 %d 页 ｜ ❌ %d 处 ｜ ⚠️  %d 处" % (len(prs.slides), n_error, n_warn))
if n_error:
    print("❌ 未通过：先修完 ❌ 再交付")
    raise SystemExit(1)
print("🎉 通过（⚠️ 为粗估：逐条修掉，或在交付说明里写明为何保留；仍需 PNG 目检）")
