# -*- coding: utf-8 -*-
"""学术 deck 交付前质检：出界 / 文字溢出 / 样板残留词 / 数字与语料交叉核验。

三道门对应三类真实返工：克隆版式后忘记收缩文本框、忘记替换样板词、
口述时说的数在页面上对不上语料。PNG 目检由 export_png.ps1 负责，本脚本只做程序化检查。
"""

import argparse
import math
import re
import sys
from pathlib import Path

from pptx import Presentation

EMU_PER_INCH = 914400.0
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
TOLERANCE_IN = 0.05
DEFAULT_SIZE_PT = 14.0

# 参考样板 deck 自带的词，出现即说明该页没替换干净。
# 两类：样板示例的学科词汇，和匿名化时填入的占位符。
LSTALE_DEFAULT = [
    "水凝胶", "粘接", "剥离", "搭接剪切", "脱粘", "PAAm", "PAA", "PDMA",
    "拓扑缠结", "速率效应", "界面失效模式一", "材料–基底界面",
    "创新成果一", "创新成果二", "创新成果三", "创新成果四",
    "张三", "李四", "20XX", "占位条目", "Placeholder", "Author A",
    "敬请各位老师批评指正",
]
# 只核验"数据型"数字：带小数点，或三位以上整数。章节号 / (a) / 01 不进核验
RE_DATA_NUMBER = re.compile(r"\d+\.\d+|\d{3,}")
# 引文年份不是数据主张，排除；真当数据用的年份由 PNG 目检兜底
RE_YEAR = re.compile(r"^(19|20)\d{2}$")
# 图注 / 引文行本就允许文字略溢出标称框，不参与溢出估算
OVERFLOW_EXEMPT_PT = 12.0


def fail(msg):
    print("❌ ERROR: " + msg)
    raise SystemExit(1)


def warn(msg):
    print("⚠️  " + msg)


def check_deck_exists(path_deck):
    if not path_deck.is_file():
        fail("deck 不存在 " + str(path_deck))
    return path_deck


def get_shape_text(shape):
    """形状（含组合内嵌）的纯文本，用换行拼接。"""
    if shape.shape_type == 6:  # msoGroup
        return "\n".join(get_shape_text(sub) for sub in shape.shapes)
    if not shape.has_text_frame:
        return ""
    return shape.text_frame.text


def get_first_size_pt(shape):
    """取形状内第一个显式字号；继承自母版时返回默认值，仅用于溢出估算。"""
    if not shape.has_text_frame:
        return DEFAULT_SIZE_PT
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                return run.font.size.pt
    return DEFAULT_SIZE_PT


def check_out_of_canvas(shape):
    """带文字的形状超出画布返回描述字符串，否则 None。

    模板的背景块、页码徽章故意出血到画布外，是设计而非缺陷，所以只查有文字的形状。
    """
    if not get_shape_text(shape).strip():
        return None
    try:
        left, top = shape.left / EMU_PER_INCH, shape.top / EMU_PER_INCH
        right = left + shape.width / EMU_PER_INCH
        bottom = top + shape.height / EMU_PER_INCH
    except TypeError:
        return None
    if (left < -TOLERANCE_IN or top < -TOLERANCE_IN
            or right > SLIDE_W_IN + TOLERANCE_IN or bottom > SLIDE_H_IN + TOLERANCE_IN):
        return "越界 (%.2f,%.2f)–(%.2f,%.2f)" % (left, top, right, bottom)
    return None


def check_text_overflow(shape):
    """按 CJK 一字一 em、西文半 em 估算行数，超出框高返回描述字符串。

    只是粗估，用于抓"克隆后塞了两倍文字"这种明显溢出，误报按 ⚠️ 处理不阻断。
    """
    text = get_shape_text(shape).strip()
    if not text or shape.shape_type == 6:
        return None
    try:
        width_in = shape.width / EMU_PER_INCH
        height_in = shape.height / EMU_PER_INCH
    except TypeError:
        return None
    size_pt = get_first_size_pt(shape)
    if size_pt <= OVERFLOW_EXEMPT_PT:
        return None
    size_in = size_pt / 72.0
    if width_in <= 0 or size_in <= 0:
        return None
    ems_per_line = width_in / size_in
    lines = 0
    for raw in text.split("\n"):
        ems = sum(1.0 if ord(ch) > 0x2E80 else 0.5 for ch in raw)
        lines += max(1, math.ceil(ems / ems_per_line))
    needed_in = lines * size_in * 1.25          # 1.25 行距
    if needed_in > height_in * 1.15:            # 留 15% 余量再报
        return "文字溢出 需%.2f in / 框高%.2f in" % (needed_in, height_in)
    return None


def check_stale_words(text, lstale):
    return [word for word in lstale if word in text]


def get_source_numbers(path_source):
    """语料里出现过的数据型数字集合；无语料时返回 None 表示跳过核验。"""
    if path_source is None:
        return None
    if not path_source.is_file():
        fail("语料文件不存在 " + str(path_source))
    return set(RE_DATA_NUMBER.findall(path_source.read_text(encoding="utf-8")))


def scan_slide(index, slide, lstale, set_source_number):
    """返回 (l错误, l警告)。"""
    lerror, lwarn = [], []
    ltext = []
    for shape in slide.shapes:
        msg = check_out_of_canvas(shape)
        if msg:
            lerror.append(msg)
        msg = check_text_overflow(shape)
        if msg:
            lwarn.append(msg)
        ltext.append(get_shape_text(shape))
    text = "\n".join(ltext)

    lhit = check_stale_words(text, lstale)
    if lhit:
        lerror.append("样板残留词 " + "、".join(lhit))

    if set_source_number is not None:
        lfound = {n for n in RE_DATA_NUMBER.findall(text) if not RE_YEAR.match(n)}
        lorphan = sorted(lfound - set_source_number)
        if lorphan:
            lerror.append("语料中查无此数 " + "、".join(lorphan))

    mark = "❌" if lerror else ("⚠️ " if lwarn else "✅")
    print("  %s 第 %d 页" % (mark, index))
    for msg in lerror + lwarn:
        print("       " + msg)
    return lerror, lwarn


def build_parser():
    parser = argparse.ArgumentParser(description="学术 deck 交付前质检")
    parser.add_argument("-deck", required=True, help="待检 .pptx")
    parser.add_argument("-source", default=None,
                        help="内容语料文本；给了才做数字交叉核验")
    parser.add_argument("-stale", default=None,
                        help="自定义残留词，逗号分隔；默认查参考样板 deck 的词")
    return parser


args = build_parser().parse_args()
sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)

path_deck = check_deck_exists(Path(args.deck).resolve())
path_source = Path(args.source).resolve() if args.source else None
lstale = [w.strip() for w in args.stale.split(",")] if args.stale else LSTALE_DEFAULT
set_source_number = get_source_numbers(path_source)

print("📁 %s" % path_deck.name)
print("📍 残留词 %d 个 ｜ 数字核验 %s"
      % (len(lstale), "开" if set_source_number is not None else "关（未给 -source）"))

prs = Presentation(str(path_deck))
n_error, n_warn = 0, 0
for index, slide in enumerate(prs.slides, 1):
    lerror, lwarn = scan_slide(index, slide, lstale, set_source_number)
    n_error += len(lerror)
    n_warn += len(lwarn)

print("📊 %d 页 ｜ ❌ %d 处 ｜ ⚠️  %d 处" % (len(prs.slides), n_error, n_warn))
if n_error:
    print("❌ 未通过：先修完 ❌ 再交付")
    raise SystemExit(1)
print("🎉 通过（⚠️ 为粗估，仍需 PNG 目检）")
